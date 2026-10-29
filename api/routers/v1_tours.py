import structlog
from fastapi import APIRouter, Depends, HTTPException, Query, Request, Header
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from typing import Annotated, Literal, Optional
from pydantic import BaseModel as _BM
import asyncpg
from api.core.aa544_tenant_pool import (
    T2_VERSION_DETAIL_READS_FLAG, T2_VERSION_READS_FLAG, T2_VERSION_WRITE_FLAG, acquire_scoped_conn,
)
from api.routers.auth import verify_tenant_jwt
from api.middleware import tenant_role as _tenant_role  # AA-979 csr check in the dependency
from api.schemas.public_api import Quota  # AA-899
from api.core.secret_compare import secrets_equal
from api.routers.admin import PLAN_LIMITS
from api.routers._billing import no_plan_usage, usage_row_to_dict
from services.acp_shared.audit_log import TenantAuditAction, write_audit_log
from api.routers._served import annotate as _annotate_served, refuse_http as _refuse_not_served
from shared import tour_liveness as _tl
from shared.tour_liveness import TourNotServedError, live_source_sql
from services.photos.sync import public_base
from api.routers._tenant_projection import (
    project_billing, project_raw, project_tour, project_version,
    tenant_status_filter, tenant_version_status,
)

logger = structlog.get_logger()
router = APIRouter(prefix="/v1/tours", tags=["B2B Tours"])
# AA-489 — separate router, plain /v1 prefix (not /v1/tours): the pre-AA-428 dead frontend
# code and this issue's own text both call this `GET /v1/quota`, not `/v1/tours/quota`. Same
# split-router-per-file pattern v1_planning.py already uses for its own slate_router.
quota_router = APIRouter(prefix="/v1", tags=["Quota"])
security = HTTPBearer()


# The tour length in days, read the way the catalog views do (mig 273 `duration_days`): the number before
# "day(s)", at most 4 digits. The old filter concatenated EVERY digit of rt.duration ("7 days 6 nights" -> 76) and
# a long string ("3 hours\n9:30am-12:30pm") overflowed INTEGER -> HTTP 500 on any duration_min / duration_max request
# (S252, AA-990 sweep). Hour-based and half/full-day products count as 1 day (they were "3 days" for "3 hours").
_DURATION_DAYS_SQL = (
    r"COALESCE(substring(lower(COALESCE(rt.duration, '')) FROM '(\d{1,4})\s*-?\s*days?')::int,"
    r" CASE WHEN lower(COALESCE(rt.duration, '')) ~ '(hour|half[ -]day|day trip|same day|full day)' THEN 1 END)"
)


def get_tenant(credentials: HTTPAuthorizationCredentials = Depends(security), request: Request = None):
    try:
        payload = verify_tenant_jwt(credentials.credentials)
    except HTTPException as exc:
        if exc.status_code == 403:      # a valid staff token is not a tenant (AA-741 pre-audit)
            raise
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    # AA-979: the dependency that accepts the token also enforces its Content studio role (csr), so a header shape
    # the middleware did not foresee cannot skip the check. request is None on a direct call (tests).
    _tenant_role.enforce(request, payload)
    return payload


def get_pool(request: Request):
    return request.app.state.pool


# AA-579: bare `GET /v1/tours` (list_tours()) removed — tàn dư kiến trúc S8/S9 (21/04/2026,
# commit a5e207d), viết cho 1 model RLS-per-tenant-row trên published_tours chưa từng thành hiện
# thực (bị đè bởi Pool/Rewrite model 2 tuần sau, 05/05/2026, cùng file). `pt.tenant_id = $1` không
# bao giờ khớp vì published_tours 100% thuộc sentinel aa_internal — route luôn trả rỗng cho mọi
# tenant thật, độc lập RLS. CloudWatch 14 ngày (/ecs/aa-cis-dev) xác nhận 0 traffic tenant thật
# (chỉ 401 từ (internal)/catalog's proxy auth mismatch — bug khác, xem AA-579). Tính năng thật đã
# được /v1/tours/pool + /v1/tours/my-versions phủ đúng kiến trúc tenant_tour_versions hiện tại.
# Xem AA-579 (Linear) cho đầy đủ bằng chứng trước khi khôi phục route này.


# ── P3-S4: Shared Pool Browse ─────────────────────────────────────────────────

# AA-859 W5 (D5a) — GET /v1/tours/pool, /my-versions and /versions/{id} run on the aa_app_user tenant pool
# behind T2_VERSION_READS_FLAG, one acquire_scoped_conn block per request, fail closed. aa_app_user reads
# the tenant's own tenant_tour_versions rows (229 RLS) and the AA master rows of published_tours /
# raw_tours (239 aa_app_user_aa_master_read). Flag off: the admin pool and the pre-AA-859 SQL, unchanged.
# Each request logs one `aa859_t2_version_read route=... used_tenant_pool=...` line. /{id}/full (tenant
# branch), /{id} and /versions/{id}/export-docx follow in D5b, behind their own flag (see _TOUR_VIA_VERSION_SQL).
_AA_MASTER_TENANT = "00000000-0000-0000-0000-000000000001"


def _log_t2_read(route: str, tenant_id, used_tenant_pool: bool, **extra) -> None:
    logger.info("aa859_t2_version_read", route=route, tenant_id=str(tenant_id),
                used_tenant_pool=used_tenant_pool, **extra)


# AA-859 MAIN exception: shared.job has no tenant column (tenant only in `payload`), is worker-hot, and
# aa_app_user has no SELECT on it after 241 (plan D2). On the tenant pool the version rows come back with
# their job_id through RLS; the statuses are read here on MAIN, after the tenant block, keyed only by those
# ids. Flag off keeps the pre-AA-859 LEFT JOIN shared.job instead.
_JOB_STATUS_SQL = "SELECT id, status FROM shared.job WHERE id = ANY($1::uuid[])"
_JOB_JOIN = "LEFT JOIN shared.job j ON j.id = ttv.job_id"


async def _job_statuses(pool, job_ids) -> dict:
    """{job_id text: status} for the given ids (None ignored); no query when there is none."""
    ids = sorted({str(j) for j in job_ids if j is not None})
    if not ids:
        return {}
    async with pool.acquire() as conn:
        rows = await conn.fetch(_JOB_STATUS_SQL, ids)
    return {str(r["id"]): r["status"] for r in rows}


# GET /v1/tours/pool. `{where}` is built from _POOL_BASE_CONDITIONS + the filters; `{tid}` is the
# placeholder number of the caller's tenant id in the `already_rewritten` EXISTS.
_POOL_BASE_CONDITIONS = (
    f"pt.tenant_id = '{_AA_MASTER_TENANT}'::uuid",
    "pt.master_status = 'active'",
    "pt.deleted_at IS NULL",
    # AA-986 L3: and its source version is live (the served rule, shared/tour_liveness.py; L2 keeps a Master off over
    # a dead source, this is the reader's half)
    live_source_sql("rt"),
)
_POOL_COUNT_SQL = """
    SELECT COUNT(*)
    FROM gold_aa_internal.published_tours pt
    LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
    {facet_join}
    {where}
"""
_POOL_REWRITTEN_TENANT_COND = "ttv.tenant_id = ${tid}::uuid"
_POOL_ROWS_SQL = """
    SELECT pt.id, pt.tour_id, pt.aa_name, pt.aa_subtitle, pt.aa_summary,
           pt.aa_highlights, pt.aa_itineraries, pt.aa_description,
           pt.seo_title, pt.seo_meta, pt.seo_keywords_used,
           pt.published_at,
           rt.country, rt.duration, rt.price_raw,
           EXISTS(
               SELECT 1 FROM gold_aa_internal.tenant_tour_versions ttv
               WHERE ttv.published_tour_id = pt.id
                 AND {rewritten_cond}
           ) AS already_rewritten,
           facet.activity_codes, facet.intensity_rank, cover.path AS cover_photo_path
    FROM gold_aa_internal.published_tours pt
    LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
    {facet_join}
    {cover_join}
    {where}
    ORDER BY {order_by}
    LIMIT ${limit} OFFSET ${offset}
"""
# AA-671 — the Marketplace facets of each pool tour, from CIS data the catalog already serves (no new column):
#   activities = the distinct graph activities of the tour's components, intensity = the highest component intensity
#   (the catalog.http.v1 facets rule, services/catalog_http/mapping.py facets()), cover = the tour's position-1 photo,
#   else a photo of its primary place (_POOL_COVER_SQL).
# Graph vocabulary: geo.tour_place CHECKs tour_place_activity_check / tour_place_intensity_level_check (graph.v1);
# the tenant sees the catalog.http.v1 intensity names (leisurely -> easy, active -> challenging), like the CRM.
# The two pools read different relations for the same rows: aa_app_user (tenant pool) holds SELECT on the catalog
# contract views and not on geo / shared.v_tour_photos; cis_app (admin pool, flag off) the reverse (ADR 0015 d3).
POOL_ACTIVITIES = ("adventure_sport", "cultural_heritage", "culinary", "local_immersion", "trekking",
                   "water_activities", "wellness_relaxation", "wildlife_nature")
POOL_INTENSITY_LEVELS = ("leisurely", "moderate", "active", "strenuous")     # increasing; rank = index + 1
POOL_INTENSITIES = ("easy", "moderate", "challenging", "strenuous")          # tenant names, same order
_INTENSITY_RANK_SQL = ("max(array_position(ARRAY['leisurely','moderate','active','strenuous']::text[], "
                       "c.intensity_level))")
_POOL_FACET_JOIN = {
    True: f"""LEFT JOIN LATERAL (
        SELECT array_agg(DISTINCT c.activity ORDER BY c.activity) AS activity_codes,
               {_INTENSITY_RANK_SQL} AS intensity_rank
          FROM catalog.v_tour_component c
         WHERE c.source_system = 'cis' AND c.tour_id = pt.tour_id
    ) facet ON true""",
    # Flag off: v_tour_component's own rows (258) on its source tables: components of the tour's source key whose
    # destination is usable. The key is written as the expression of raw_tours.source_key (258,
    # COALESCE(source_group_id, tour_id)::text, a stored generated column on Dev), like 272, so the SQL does not
    # depend on the stored column (the integration replay skips 258).
    False: f"""LEFT JOIN LATERAL (
        SELECT array_agg(DISTINCT c.activity ORDER BY c.activity) AS activity_codes,
               {_INTENSITY_RANK_SQL} AS intensity_rank
          FROM geo.tour_place c
          JOIN shared.destinations d ON d.id = c.destination_id
                                    AND d.geo_status <> 'suspect' AND d.duplicate_of IS NULL
         WHERE c.source_id = (COALESCE(rt.source_group_id, rt.tour_id))::text
    ) facet ON true""",
}
# AA-671 S257 — cover fallback. Only the tours with a photo folder of their own have a position-1 tour photo (S256:
# 162 of 892 pool tours), so a tour without one takes a photo of its primary place: among the usable places of its
# graph components (the facet join's rows) that have an AA photo, the one the tour spends the most days in, then the
# earliest day, then the destination id (deterministic). Within that place the photo is picked by
# md5(tour id || path), stable for a tour but spread across tours, so twenty Hanoi tours on one grid page do not all
# show the same picture (lesson S253, duplicate covers). Destination photos come from the same photo read views
# (matched, served rows: every one has the 600 / 960 / 1600 px renditions), so the URLs are built the same way.
_POOL_COVER_SQL = """LEFT JOIN LATERAL (
        SELECT COALESCE(
            (SELECT ph.path FROM {tour_photo} ph
              WHERE {tour_photo_cond} ph.tour_id = pt.tour_id AND ph.position = 1
              LIMIT 1),
            (SELECT place_photo.path
               FROM (SELECT c.destination_id, count(DISTINCT c.day_index) AS days, min(c.day_index) AS first_day
                       FROM {components}
                      GROUP BY c.destination_id) place
              CROSS JOIN LATERAL (
                    SELECT dp.path FROM {destination_photo} dp
                     WHERE dp.destination_id = place.destination_id
                     ORDER BY md5(pt.tour_id::text || dp.path)
                     LIMIT 1) place_photo
              ORDER BY place.days DESC, place.first_day, place.destination_id
              LIMIT 1)
        ) AS path
    ) cover ON true"""
_POOL_COVER_JOIN = {
    True: _POOL_COVER_SQL.format(
        tour_photo="catalog.v_tour_photo", tour_photo_cond="ph.source_system = 'cis' AND",
        components="catalog.v_tour_component c WHERE c.source_system = 'cis' AND c.tour_id = pt.tour_id",
        destination_photo="catalog.v_destination_photo"),
    # Flag off: the facet join's source rows (geo.tour_place on a usable destination, by the tour's source key)
    False: _POOL_COVER_SQL.format(
        tour_photo="shared.v_tour_photos", tour_photo_cond="",
        components="""geo.tour_place c
                       JOIN shared.destinations d ON d.id = c.destination_id
                                                 AND d.geo_status <> 'suspect' AND d.duplicate_of IS NULL
                      WHERE c.source_id = (COALESCE(rt.source_group_id, rt.tour_id))::text""",
        destination_photo="shared.v_destination_photos"),
}


def _pool_facets(row: dict) -> dict:
    """The facet columns of one pool row -> the tenant fields (activities, intensity, cover image URLs)."""
    codes = row.pop("activity_codes", None) or []
    rank = row.pop("intensity_rank", None)
    path = row.pop("cover_photo_path", None)
    row["activities"] = [c for c in codes if c in POOL_ACTIVITIES]
    row["intensity"] = POOL_INTENSITIES[rank - 1] if isinstance(rank, int) and 1 <= rank <= 4 else None
    base = public_base() + path if isinstance(path, str) and path.startswith("/") else None
    row["cover_image_url"] = f"{base}?size=small" if base else None
    row["cover_image_large_url"] = f"{base}?size=large" if base else None
    return row


# Countries for the filter dropdown: same master_status/deleted_at gate as the listing, so a
# trashed/deactivated tour's country doesn't show as a filterable option.
_POOL_COUNTRIES_SQL = f"""
    SELECT DISTINCT rt.country
    FROM gold_aa_internal.published_tours pt
    LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
    WHERE pt.tenant_id = '{_AA_MASTER_TENANT}'::uuid
      AND pt.master_status = 'active'
      AND pt.deleted_at IS NULL
      AND {live_source_sql("rt")}
      AND rt.country IS NOT NULL
    ORDER BY rt.country
"""


@router.get("/pool")
async def browse_pool(
    request: Request,
    tenant=Depends(get_tenant),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    country: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    duration_min: Optional[int] = Query(None, ge=1),
    duration_max: Optional[int] = Query(None),
    sort: str = Query("newest"),
    # Annotated, so a direct call (tests, other routers) that leaves them out gets None, not a Query object
    activity: Annotated[Optional[Literal[POOL_ACTIVITIES]], Query()] = None,
    intensity: Annotated[Optional[Literal[POOL_INTENSITIES]], Query()] = None,
    added_within_days: Annotated[Optional[int], Query(ge=1, le=365)] = None,
):
    """Browse AA shared pool (published_tours from aa_internal).

    AA-671: each row also carries `activities` (graph activity codes), `intensity` (easy | moderate | challenging |
    strenuous, the highest of the tour's components, or null) and `cover_image_url` / `cover_image_large_url` (the
    tour's first photo on the public photo route, 600 / 1600 px; else a photo of the place the tour spends the most
    days in; null when neither exists). Filters `activity`, `intensity` (exact
    level) and `added_within_days` (published in the last N days); the answer's `facets` lists the accepted values.

    AA-796: no A1 quality score on the tenant side — the `min_quality` filter, the
    `sort=quality` order and the `quality_score` column are gone. An old client still sending
    `min_quality` / `sort=quality` gets the unfiltered, newest-first list (FastAPI ignores
    unknown query params).
    """
    tenant_id = tenant["sub"]
    # The listing is the AA master tenant's published tours (the shared pool itself), read by every
    # active tenant. AA-859 W5: on the tenant pool aa_app_user reads them through 239's master-read
    # policies, and the `already_rewritten` EXISTS sees only the caller's own tenant_tour_versions (229
    # RLS) even if its tenant condition were wrong. Flag off: the admin pool, as before (AA-544 Round 2
    # decision (a)).
    offset = (page - 1) * page_size

    # AA-441 (bug #6, AA-438-00-SUMMARY #8): was missing the master_status/deleted_at gate that
    # acp_contract.v_trip_registry already applies (migrations 078/083/090) — a tour an admin
    # trashes/deactivates in Master Content stayed fully visible and rewritable to every tenant
    # here. Same condition v_trip_registry uses: master_status='active' AND deleted_at IS NULL.
    conditions = list(_POOL_BASE_CONDITIONS)
    params: list = []

    if country:
        params.append(country)
        conditions.append(f"LOWER(rt.country) = LOWER(${len(params)})")
    if search:
        params.append(f"%{search}%")
        conditions.append(f"pt.aa_name ILIKE ${len(params)}")
    if duration_min is not None:
        params.append(duration_min)
        conditions.append(
            f"{_DURATION_DAYS_SQL} >= ${len(params)}"
        )
    if duration_max is not None:
        params.append(duration_max)
        conditions.append(
            f"{_DURATION_DAYS_SQL} <= ${len(params)}"
        )

    # AA-671 facet filters (the facet join is added to the count only when one of them is set)
    if activity is not None:
        params.append(activity)
        conditions.append(f"${len(params)}::text = ANY(facet.activity_codes)")
    if intensity is not None:
        params.append(POOL_INTENSITIES.index(intensity) + 1)
        conditions.append(f"facet.intensity_rank = ${len(params)}::int")
    if added_within_days is not None:
        params.append(added_within_days)
        conditions.append(f"pt.published_at >= now() - make_interval(days => ${len(params)}::int)")
    count_needs_facets = activity is not None or intensity is not None

    order_by = "pt.published_at DESC"  # AA-796 — the only tenant sort; `sort` kept for old clients
    where = "WHERE " + " AND ".join(conditions)
    # tenant_id added as last param for the already_rewritten subquery
    params_paged = params + [page_size, offset, tenant_id]

    async with acquire_scoped_conn(
        request, tenant_id, T2_VERSION_READS_FLAG, fail_closed=True,
    ) as (conn, used_tenant_pool):
        facet_join = _POOL_FACET_JOIN[used_tenant_pool]
        rows_sql = _POOL_ROWS_SQL.format(
            rewritten_cond=_POOL_REWRITTEN_TENANT_COND.format(tid=len(params) + 3),
            facet_join=facet_join, cover_join=_POOL_COVER_JOIN[used_tenant_pool],
            where=where, order_by=order_by, limit=len(params) + 1, offset=len(params) + 2)
        count_sql = _POOL_COUNT_SQL.format(facet_join=facet_join if count_needs_facets else "", where=where)
        total = await conn.fetchval(count_sql, *params)
        rows = await conn.fetch(rows_sql, *params_paged)
        countries = await conn.fetch(_POOL_COUNTRIES_SQL)
        data = await _annotate_served(conn, [_pool_facets(project_tour(r)) for r in rows])   # AA-986 L3
    _log_t2_read("pool", tenant_id, used_tenant_pool, count=len(rows))

    return {
        "data": data,
        "pagination": {"page": page, "page_size": page_size,
                       "total": total, "pages": -(-total // page_size)},
        "countries": [r["country"] for r in countries],
        "facets": {"activities": list(POOL_ACTIVITIES), "intensities": list(POOL_INTENSITIES)},
        "tenant_id": tenant_id,
    }


# ── AA-489: rewrite quota (shared helpers + GET /v1/quota) ────────────────────
# STEP0 found nothing enforces a tenant's monthly rewrite count: PLAN_LIMITS.tours_per_month
# (admin.py) was defined but only ever displayed, never checked; raw_tours.rewrite_count
# (migration 033) is dead, 0 callers; rate_limit_middleware throttles requests/minute across
# ALL /v1/*, not a monthly count. Deliberately NOT reusing acp_quota_ledger — that's ACPv1
# S2/S3/S4 quota, dead since 13/07/2026. New table: shared.tenant_rewrite_usage (migration
# 142), one row per (tenant_id, year_month).
#
# Business decisions — confirmed by Nghiệp 25/09/2026 (AA-640), replacing AA-489's defaults:
#   * Limit = shared.membership_plans.tours_quota_monthly for the tenant's plan — the SAME number
#     the portal shows and billing (v_tenant_monthly_usage overage) uses. PLAN_LIMITS in admin.py
#     used to carry a second, different set (growth 500 vs 200 shown) — it now only holds RPM.
#   * Reset = calendar month, not rolling 30d.
#   * NO hard block while the product is in trial: a rewrite past the quota is allowed and billed
#     as overage (overage_rate_usd_per_tour, already computed by v_tenant_monthly_usage). Logged as
#     `rewrite_over_quota` so it stays visible.


def _current_year_month_and_reset() -> tuple:
    """(year_month 'YYYY-MM' string, first-of-next-month date) in UTC."""
    import datetime as _dt_quota
    now_utc = _dt_quota.datetime.now(_dt_quota.timezone.utc)
    next_month = (now_utc.replace(day=1) + _dt_quota.timedelta(days=32)).replace(day=1)
    return now_utc.strftime("%Y-%m"), next_month


async def _get_tenant_plan_limit(conn, tenant_id: str) -> tuple:
    """(plan_tier str, tours_quota_monthly int) — plan read live from shared.tenants, not the
    JWT's own plan_tier claim (same rationale AA-432 established for rate_limit_rpm/is_active: a
    plan change shouldn't take up to 24h, the JWT's TTL, to take effect). The quota comes from
    shared.membership_plans (AA-640); a plan with no row falls back to the starter row."""
    row = await conn.fetchrow("""
        SELECT t.plan_tier::text AS plan,
               COALESCE(mp.tours_quota_monthly,
                        (SELECT tours_quota_monthly FROM shared.membership_plans
                         WHERE plan_name = 'starter')) AS quota
        FROM shared.tenants t
        LEFT JOIN shared.membership_plans mp ON mp.plan_name = t.plan_tier::text
        WHERE t.tenant_id = $1::uuid
    """, tenant_id)
    if not row:
        return "starter", 0
    return str(row["plan"]), int(row["quota"] or 0)


async def _check_and_consume_rewrite_quota(conn, tenant_id: str) -> None:
    """Atomically increments this month's rewrite count. AA-640: past the plan quota the rewrite
    is still allowed (billed as overage) — logged, never blocked, while the product is in trial."""
    plan, limit = await _get_tenant_plan_limit(conn, tenant_id)
    year_month, _next_month = _current_year_month_and_reset()
    used = await conn.fetchval("""
        INSERT INTO shared.tenant_rewrite_usage (tenant_id, year_month, rewrite_count)
        VALUES ($1::uuid, $2, 1)
        ON CONFLICT (tenant_id, year_month)
        DO UPDATE SET rewrite_count = shared.tenant_rewrite_usage.rewrite_count + 1,
                      updated_at = NOW()
        RETURNING rewrite_count
    """, tenant_id, year_month)
    if used > limit:
        logger.info("rewrite_over_quota", tenant_id=tenant_id, plan=plan, used=used, quota=limit,
                    year_month=year_month)


# Real replacement for the endpoint CatalogTab.tsx used to call before AA-428 deleted that
# dead code (STEP0 there found this route never existed). Read-only — does NOT increment
# shared.tenant_rewrite_usage (only _check_and_consume_rewrite_quota, called from
# trigger_rewrite() below, does that, on an actual attempt). `rewrites_remaining` kept as the
# field name the old removed UI already expected, so a future UI rebuild (the issue's own "nếu
# muốn") can reuse this shape as-is.
@quota_router.get("/quota", response_model=Quota)  # AA-899: model
async def get_quota(request: Request, tenant=Depends(get_tenant)):
    tenant_id = tenant["sub"]
    pool = request.app.state.pool
    year_month, next_month = _current_year_month_and_reset()

    async with pool.acquire() as conn:
        plan, limit = await _get_tenant_plan_limit(conn, tenant_id)
        used = await conn.fetchval("""
            SELECT rewrite_count FROM shared.tenant_rewrite_usage
            WHERE tenant_id = $1::uuid AND year_month = $2
        """, tenant_id, year_month) or 0

    return {
        "plan_tier":           plan,
        "tours_per_month":     limit,
        "rewrites_used":       used,
        "rewrites_remaining":  max(limit - used, 0),
        "resets_at":           next_month.strftime("%Y-%m-%d"),
    }


# ── AA-496: GET /v1/billing — tenant-scoped self-service billing view ─────────
# STEP0 found: `/portal/*`'s layout.tsx + DashboardTab.tsx both `fetch("/api/admin/billing")`
# unconditionally on every page load. That proxy (`/api/admin/[...path]/route.ts`) calls
# `requireAdmin()` first — a real, independent admin-JWT check — before ever attaching
# X-Admin-Secret. A tenant portal session is never an admin session, so this 401s at the
# Next.js proxy layer on literally every real tenant, on every page, always — confirmed by
# reading both call sites and the backend `GET /admin/billing` (admin_pipeline.py): it takes
# an arbitrary `?tenant_id=` query param with NO scoping to "the caller's own tenant" (by
# design — it's an admin looking-glass over ANY tenant, defaulting to aa_internal), so it
# could never safely be exposed to a tenant session even with a header change alone.
# This is possibility #1 from the issue text ("gọi sai chỗ"), not a missing-feature gap:
# `/v1/quota` (AA-489, just above) already proves the tenant-scoped-via-JWT pattern this
# needed. Real fix: a tenant-scoped sibling endpoint, JWT-derived tenant_id (never trusts a
# query param), reusing the exact same v_tenant_monthly_usage view + shape the admin view
# already returns (so BillingTab.tsx/DashboardTab.tsx need zero shape changes) — frontend
# repointed from /api/admin/billing to /api/tenant/v1/billing (see those 2 files' own diffs).
@quota_router.get("/billing")
async def get_my_billing(request: Request, tenant=Depends(get_tenant)):
    tenant_id = tenant["sub"]
    pool = request.app.state.pool
    async with pool.acquire() as conn:
        row = await conn.fetchrow("""
            SELECT
                v.tenant_name, v.plan_tier, v.billing_month,
                v.tours_quota_monthly, v.api_calls_quota_monthly,
                v.price_usd_monthly,
                v.tours_rewritten, v.api_calls_used,
                v.quota_tours_pct, v.quota_calls_pct,
                v.tours_overage, v.overage_usd,
                v.overage_rate_usd_per_tour
            FROM shared.v_tenant_monthly_usage v
            WHERE v.tenant_id = $1::uuid
        """, tenant_id)

        activity = await conn.fetch("""
            SELECT ttv.id, ttv.created_at, ttv.status, ttv.edit_source,
                   pt.aa_name, rt.country
            FROM gold_aa_internal.tenant_tour_versions ttv
            JOIN gold_aa_internal.published_tours pt ON pt.id = ttv.published_tour_id
            LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
            WHERE ttv.tenant_id = $1::uuid
            ORDER BY ttv.created_at DESC LIMIT 5
        """, tenant_id)

        # AA-636: the portal hardcoded "300 RPM", "20,000 API calls" and a Growth-is-current plan
        # table for every tenant. Serve the tenant's real rate limit and the sellable plans
        # (shared.membership_plans, the same table v_tenant_monthly_usage reads) instead.
        rate_limit_rpm = await conn.fetchval(
            "SELECT rate_limit_rpm FROM shared.tenants WHERE tenant_id = $1::uuid", tenant_id
        )
        plan_rows = await conn.fetch("""
            SELECT plan_name, tours_quota_monthly, api_calls_quota_monthly, price_usd_monthly
            FROM shared.membership_plans
            WHERE is_active AND plan_name <> 'internal'
            ORDER BY (price_usd_monthly IS NULL OR price_usd_monthly <= 0), price_usd_monthly, tours_quota_monthly
        """)
        # AA-638 — per-day API calls + tour rewrites for the current month (Dashboard sparkline).
        daily_rows = await conn.fetch("""
            WITH days AS (
                SELECT generate_series(date_trunc('month', now()), date_trunc('day', now()),
                                       interval '1 day') AS day
            ),
            api AS (
                SELECT date_trunc('day', called_at) AS day, COUNT(*) AS n
                FROM shared.tenant_api_usage
                WHERE tenant_id = $1::uuid AND called_at >= date_trunc('month', now())
                GROUP BY 1
            ),
            rw AS (
                SELECT date_trunc('day', created_at) AS day, COUNT(*) AS n
                FROM gold_aa_internal.tenant_tour_versions
                WHERE tenant_id = $1::uuid AND created_at >= date_trunc('month', now())
                GROUP BY 1
            )
            SELECT days.day::date AS day, COALESCE(api.n, 0) AS api_calls, COALESCE(rw.n, 0) AS rewrites
            FROM days
            LEFT JOIN api ON api.day = days.day
            LEFT JOIN rw ON rw.day = days.day
            ORDER BY days.day
        """, tenant_id)
    daily = [
        {"day": str(d["day"]), "api_calls": int(d["api_calls"]), "rewrites": int(d["rewrites"])}
        for d in daily_rows
    ]
    plans = [
        {
            "plan_name": p["plan_name"],
            "tours_quota_monthly": p["tours_quota_monthly"],
            "api_calls_quota_monthly": p["api_calls_quota_monthly"],
            # enterprise is stored as 0.00 = negotiated, not free -> served as None ("Custom")
            "price_usd_monthly": (
                float(p["price_usd_monthly"]) if p["price_usd_monthly"] and p["price_usd_monthly"] > 0 else None
            ),
            "rate_limit_rpm": PLAN_LIMITS.get(p["plan_name"], {}).get("rpm"),
        }
        for p in plan_rows
    ]

    # AA-798: no usage row -> no plan, nothing counted. It used to invent a Starter plan
    # (fixed price / 50 tours / 5000 calls); the portal now shows "No plan attached" instead.
    usage = usage_row_to_dict(project_billing(row)) if row else no_plan_usage()
    return {
        **usage,
        "rate_limit_rpm": rate_limit_rpm,
        "plans": plans,
        "daily": daily,
        "activity": [
            {
                "id": str(a["id"]),
                "created_at": a["created_at"].isoformat(),
                "status": tenant_version_status(a["status"]),  # AA-796
                "edit_source": a["edit_source"],
                "tour_name": a["aa_name"],
                "country": a["country"],
            }
            for a in activity
        ],
    }


# ── P3-S4: Trigger Rewrite ────────────────────────────────────────────────────


class RewriteRequest(_BM):
    rewrite_language: str = "en-US"
    seo_mode: str = "standard"
    custom_notes: Optional[str] = None


@router.post("/pool/{published_tour_id}/rewrite")
async def trigger_rewrite(
    published_tour_id: str,
    body: RewriteRequest,
    request: Request,
    tenant=Depends(get_tenant),
):
    """Trigger tenant rewrite of a published tour.
    Creates tenant_tour_versions record (status=pending) and calls pipeline.
    """
    tenant_id = tenant["sub"]
    pool = request.app.state.pool

    # AA-741 pre-audit: validate first, charge last, all in ONE transaction. Before, the quota was
    # consumed (autocommit) before the tour was even looked up, so a 404, an inactive Master or a
    # double-click each billed a rewrite; and the tour was not checked against the same
    # active-Master gate the pool list applies (`list_pool`: master_status='active' AND
    # deleted_at IS NULL), so a tenant could rewrite a withdrawn tour by id.
    async with pool.acquire() as conn, conn.transaction():
        pt = await conn.fetchrow("""
            SELECT pt.id, pt.tour_id, pt.aa_name, pt.aa_subtitle,
                   pt.aa_summary, pt.aa_description, pt.aa_highlights,
                   pt.aa_itineraries, pt.seo_title, pt.seo_meta,
                   pt.seo_keywords_used,
                   rt.country, rt.duration
            FROM gold_aa_internal.published_tours pt
            LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
            WHERE pt.id = $1::uuid
              AND pt.master_status = 'active' AND pt.deleted_at IS NULL
              AND rt.source_status = 'active' AND rt.deleted_at IS NULL
        """, published_tour_id)
        if not pt:
            # AA-986 L3 (ADR 0015 decision 5): a tour AA no longer serves (retired) or has paused answers 409 with a
            # plain sentence, not the generic 404; nothing is charged (the quota is consumed below).
            tour_of = await conn.fetchval(
                "SELECT tour_id FROM gold_aa_internal.published_tours WHERE id = $1::uuid", published_tour_id)
            try:
                await _tl.ensure_served(conn, tour_of, "start a new rewrite")
            except TourNotServedError as exc:
                raise _refuse_not_served(exc) from exc
            raise HTTPException(status_code=404, detail="Tour not found in pool")

        # One rewrite in flight per (tenant, tour): serialise concurrent clicks on this pair, then
        # refuse a second pending version instead of billing and writing it twice.
        await conn.execute("SELECT pg_advisory_xact_lock(hashtext($1))",
                           f"t2_rewrite:{tenant_id}:{published_tour_id}")
        in_flight = await conn.fetchval("""
            SELECT id FROM gold_aa_internal.tenant_tour_versions
            WHERE tenant_id = $1::uuid AND published_tour_id = $2::uuid AND status = 'pending'
            ORDER BY created_at DESC LIMIT 1
        """, tenant_id, published_tour_id)
        if in_flight:
            raise HTTPException(status_code=409,
                                detail="A rewrite of this tour is already in progress")

        # AA-489/AA-640 — count this rewrite against the monthly plan quota. Over-quota rewrites
        # are allowed and billed as overage (trial phase) — see the helper's comment block.
        await _check_and_consume_rewrite_quota(conn, tenant_id)

        # Get next version number
        next_ver = await conn.fetchval("""
            SELECT COALESCE(MAX(version_number), 0) + 1
            FROM gold_aa_internal.tenant_tour_versions
            WHERE tenant_id = $1::uuid AND published_tour_id = $2::uuid
        """, tenant_id, published_tour_id)

        # Create pending version record
        import json as _json
        version_id = await conn.fetchval("""
            INSERT INTO gold_aa_internal.tenant_tour_versions
                (tenant_id, published_tour_id, version_number,
                 rewritten_content, status, edit_source,
                 rewrite_language, seo_mode)
            VALUES ($1::uuid, $2::uuid, $3,
                    $4::jsonb, 'pending', 'ai_generated',
                    $5, $6)
            RETURNING id
        """,
            tenant_id, published_tour_id, next_ver,  # noqa: E128
            _json.dumps({  # noqa: E128
                "name": pt["aa_name"], "summary": pt["aa_summary"],
                "status": "generating",
            }),
            body.rewrite_language, body.seo_mode)

        # AA-559 — semantic tenant-activity log, same connection as the pending-version INSERT
        # above (T1 Browse Pool "rewrite" action).
        await write_audit_log(
            conn, tenant_id=str(tenant_id), actor=f"tenant:{tenant_id}",
            action=TenantAuditAction.TOUR_REWRITE_TRIGGERED, resource_type="tenant_tour_version",
            resource_id=str(version_id),
            details={
                "published_tour_id": published_tour_id, "version_number": next_ver,
                "tour_name": pt["aa_name"], "rewrite_language": body.rewrite_language,
            },
        )

    # AA-652 — the rewrite (SEO lookup, _rewrite_tour, T3 QA gate, save, live progress) runs as
    # a durable `t2_rewrite` job (services/jobs/t2_rewrite_job.py). Before, it was an in-process
    # asyncio task: a deploy/restart killed it and the version stayed 'pending' ("Writing…")
    # forever. The idempotency key ties one job to one version.
    from services.jobs.t2_rewrite_job import KIND as _T2_KIND
    from shared.jobs.registry import enqueue as _enqueue_job
    job_id, _ = await _enqueue_job(
        pool, _T2_KIND,
        {"version_id": str(version_id), "tenant_id": str(tenant_id),
         "published_tour_id": str(published_tour_id), "rewrite_language": body.rewrite_language},
        idempotency_key=f"{_T2_KIND}:{version_id}", created_by=f"tenant:{tenant_id}",
    )
    async with pool.acquire() as conn:
        await conn.execute(
            "UPDATE gold_aa_internal.tenant_tour_versions SET job_id = $2::uuid WHERE id = $1::uuid",
            version_id, job_id)

    return {
        "version_id": str(version_id),
        "published_tour_id": published_tour_id,
        "version_number": next_ver,
        "status": "pending",
        "job_id": job_id,
        "message": "Rewrite started — check My Catalog in ~30 seconds for results",
    }


# ── P3-S4: My Versions (Tenant Catalog) ──────────────────────────────────────

# GET /v1/tours/my-versions. `{where}` starts with _MY_VERSIONS_TENANT_COND; `{job_col}` / `{job_join}`:
# flag off, the job status joined from shared.job (pre-AA-859); tenant pool, the job id (_MY_VERSIONS_JOB_SPLIT).
_MY_VERSIONS_TENANT_COND = "ttv.tenant_id = $1::uuid"
_MY_VERSIONS_COUNT_SQL = """
    SELECT COUNT(*)
    FROM gold_aa_internal.tenant_tour_versions ttv
    JOIN gold_aa_internal.published_tours pt ON pt.id = ttv.published_tour_id
    LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
    {where}
"""
_MY_VERSIONS_ROWS_SQL = """
    SELECT ttv.id, ttv.version_number, ttv.status,
           ttv.edit_source, ttv.rewrite_language, ttv.created_at,
           ttv.rewritten_content,
           pt.id AS published_tour_id, pt.tour_id, pt.aa_name,
           pt.updated_at AS master_updated_at,  -- AA-990: master_newer is derived from it (project_version)
           rt.country, rt.duration,
           -- AA-652 — the rewrite job's status (queued/running/succeeded/failed/...),
           -- NULL for manual edits and pre-AA-652 rows. No job error text (tenant-safe).
           {job_col}
    FROM gold_aa_internal.tenant_tour_versions ttv
    JOIN gold_aa_internal.published_tours pt ON pt.id = ttv.published_tour_id
    LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
    {job_join}
    {where}
    ORDER BY ttv.created_at DESC, ttv.id DESC
    LIMIT ${limit} OFFSET ${offset}
"""
_MY_VERSIONS_JOB_JOINED = {"job_col": "j.status AS job_status", "job_join": _JOB_JOIN}
_MY_VERSIONS_JOB_SPLIT = {"job_col": "ttv.job_id AS _job_id, NULL::text AS job_status", "job_join": ""}


async def _with_job_status(pool, rows, job_id_key: str) -> list:
    """Tenant-pool rows -> dicts with `job_status` filled from shared.job on MAIN (D2 exception, see
    _JOB_STATUS_SQL). `_job_id` (my-versions' helper column) is dropped; same keys as the joined form."""
    statuses = await _job_statuses(pool, [r[job_id_key] for r in rows])
    out = []
    for r in rows:
        d = dict(r)
        job_id = d[job_id_key]
        if job_id_key == "_job_id":
            del d["_job_id"]  # my-versions' helper column, not part of the response
        d["job_status"] = statuses.get(str(job_id)) if job_id is not None else None  # keeps the key's place
        out.append(d)
    return out


@router.get("/my-versions")
async def list_my_versions(
    request: Request,
    tenant=Depends(get_tenant),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    status: Optional[str] = Query(None),
    country: Optional[str] = Query(None),
):
    """List tenant's own rewritten versions.

    AA-990: each row also carries `master_updated_at` (when Adventure Asia last changed the Master text of the
    tour) and `master_newer` (true when that is later than the version's `created_at`). Read-only, informational:
    nothing in the tenant's version is changed."""
    tenant_id = tenant["sub"]
    offset = (page - 1) * page_size

    conditions = [_MY_VERSIONS_TENANT_COND]
    params: list = [tenant_id]

    if status:
        # AA-796 — the tenant only ever sees ai_generated; asking for it (or for the stored
        # needs_review) matches both stored values.
        params.append(tenant_status_filter(status))
        conditions.append(f"ttv.status = ANY(${len(params)}::text[])")
    if country:
        params.append(country)
        conditions.append(f"LOWER(rt.country) = LOWER(${len(params)})")

    where = "WHERE " + " AND ".join(conditions)

    async with acquire_scoped_conn(
        request, tenant_id, T2_VERSION_READS_FLAG, fail_closed=True,
    ) as (conn, used_tenant_pool):
        total = await conn.fetchval(_MY_VERSIONS_COUNT_SQL.format(where=where), *params)
        job = _MY_VERSIONS_JOB_SPLIT if used_tenant_pool else _MY_VERSIONS_JOB_JOINED
        rows = await conn.fetch(_MY_VERSIONS_ROWS_SQL.format(
            where=where, limit=len(params) + 1, offset=len(params) + 2, **job), *params, page_size, offset)
        served = await _annotate_served(conn, [{"tour_id": r.get("tour_id")} for r in rows])   # AA-986 L3
    _log_t2_read("my_versions", tenant_id, used_tenant_pool, count=len(rows))
    if used_tenant_pool:  # after the tenant block: the tenant connection is released before MAIN
        rows = await _with_job_status(request.app.state.pool, rows, "_job_id")

    return {
        "data": [{**project_version(r), **{k: v for k, v in sv.items() if k != "tour_id"}}
                 for r, sv in zip(rows, served)],
        "pagination": {"page": page, "page_size": page_size,
                       "total": total, "pages": -(-total // page_size)},
    }


@router.post("/versions/{version_id}/retry")
async def retry_rewrite(version_id: str, request: Request, tenant=Depends(get_tenant)):
    """AA-652 — re-run a rewrite whose job failed (the portal's Retry button). Re-queues the same
    job with fresh attempts; no second quota charge (the rewrite was already counted)."""
    from shared.jobs import queue as job_queue

    tenant_id = tenant["sub"]
    pool = request.app.state.pool
    async with pool.acquire() as conn:
        row = await conn.fetchrow("""
            SELECT ttv.id, ttv.status, ttv.job_id, pt.tour_id
            FROM gold_aa_internal.tenant_tour_versions ttv
            LEFT JOIN gold_aa_internal.published_tours pt ON pt.id = ttv.published_tour_id
            WHERE ttv.id = $1::uuid AND ttv.tenant_id = $2::uuid
        """, version_id, tenant_id)
        if row and row["status"] == "failed" and row["job_id"] is not None:
            # AA-986 L3: no retry (LLM spend) on a tour AA no longer serves or has paused
            try:
                await _tl.ensure_served(conn, row.get("tour_id"), "retry this rewrite")
            except TourNotServedError as exc:
                raise _refuse_not_served(exc) from exc
    if not row:
        raise HTTPException(status_code=404, detail="Version not found")
    if row["status"] != "failed" or row["job_id"] is None:
        raise HTTPException(status_code=409, detail="Only a failed AI rewrite can be retried")
    # AA-741 pre-audit: flip the version to pending BEFORE re-queueing. The job skips any version
    # that is not 'pending', so a worker that claimed the re-queued job before this UPDATE landed
    # finished it as {"skipped": true} and the version stayed failed.
    async with pool.acquire() as conn:
        flipped = await conn.fetchval(
            "UPDATE gold_aa_internal.tenant_tour_versions SET status = 'pending' "
            "WHERE id = $1::uuid AND status = 'failed' RETURNING id", version_id)
    if not flipped:
        raise HTTPException(status_code=409, detail="This rewrite is already running")
    if not await job_queue.retry(pool, str(row["job_id"])):
        async with pool.acquire() as conn:
            await conn.execute(
                "UPDATE gold_aa_internal.tenant_tour_versions SET status = 'failed' "
                "WHERE id = $1::uuid AND status = 'pending'", version_id)
        raise HTTPException(status_code=409, detail="This rewrite is already running")
    async with pool.acquire() as conn:
        await write_audit_log(
            conn, tenant_id=str(tenant_id), actor=f"tenant:{tenant_id}",
            action=TenantAuditAction.TOUR_REWRITE_TRIGGERED, resource_type="tenant_tour_version",
            resource_id=str(version_id), details={"retry": True, "job_id": str(row["job_id"])},
        )
    return {"version_id": version_id, "status": "pending", "job_id": str(row["job_id"])}


# ── AA-859 D5b: T2 detail reads on the tenant pool (T2_VERSION_DETAIL_READS_FLAG) ─────────────────
# GET /v1/tours/{id}/full (tenant-JWT branch only), GET /v1/tours/{id} and GET /v1/tours/versions/{id}/export-docx,
# one acquire_scoped_conn block per request, fail closed. aa_app_user reads the tour through the caller's own
# tenant_tour_versions rows (229 RLS) and the AA master rows of published_tours / raw_tours / seo_context (239
# aa_app_user_aa_master_read). Flag off: the admin pool and the pre-AA-859 SQL text, unchanged. Its own flag,
# not T2_VERSION_READS_FLAG (that one is ON since S240: a route added under it would move with no flip of its
# own, and reverting one would revert the other). One `aa859_t2_version_detail_read route=full|tour|export_docx
# used_tenant_pool=...` line per request. The /full admin branch (X-Admin-Secret) never reads the flag: staff are
# outside the tenant boundary (plan D2), it stays on the admin pool. SQL in module constants so tests can mutate
# the tenant conditions.
_TOUR_VIA_VERSION_TENANT_COND = "ttv.tenant_id = $2::uuid"
_TOUR_VIA_VERSION_SQL = """
            SELECT pt.* FROM gold_aa_internal.published_tours pt
            JOIN gold_aa_internal.tenant_tour_versions ttv
                ON ttv.published_tour_id = pt.id
            WHERE pt.id = $1::uuid AND {tenant_cond}
            LIMIT 1
        """
_TOUR_RAW_SQL = """
            SELECT * FROM silver_aa_internal.raw_tours
            WHERE tour_id = $1::uuid
        """
_TOUR_SEO_SQL = """
            SELECT top_keywords, keyword_search, keyword_ideas, fetched_at
            FROM silver_aa_internal.seo_context
            WHERE tour_id = $1::uuid
            ORDER BY fetched_at DESC LIMIT 1
        """


def _log_t2_detail_read(route: str, tenant_id, used_tenant_pool: bool, **extra) -> None:
    logger.info("aa859_t2_version_detail_read", route=route, tenant_id=str(tenant_id),
                used_tenant_pool=used_tenant_pool, **extra)


def _safe_row(row):
    from uuid import UUID
    from decimal import Decimal
    if not row: return {}
    d = dict(row)
    for k, v in d.items():
        if isinstance(v, UUID):
            d[k] = str(v)
        elif isinstance(v, Decimal):
            d[k] = float(v)
        elif hasattr(v, 'isoformat'):
            d[k] = v.isoformat()
    return d


@router.get("/{tour_id}/full")
async def get_tour_full(
    tour_id: str,
    request: Request,
    x_admin_secret: Optional[str] = Header(None),
    credentials: Optional[HTTPAuthorizationCredentials] = Depends(HTTPBearer(auto_error=False)),
):
    """Full before/after data for catalog review panel."""
    import os
    ADMIN_SECRET = os.environ.get("ADMIN_SECRET", "")
    is_admin = bool(ADMIN_SECRET and secrets_equal(x_admin_secret, ADMIN_SECRET))
    if is_admin:
        tenant_id = "00000000-0000-0000-0000-000000000001"
    elif credentials:
        try:
            payload = verify_tenant_jwt(credentials.credentials)
            tenant_id = payload["sub"]
        except Exception:
            raise HTTPException(status_code=401, detail="Invalid or expired token")
    else:
        raise HTTPException(status_code=401, detail="Not authenticated")

    if is_admin:
        # AA-859 MAIN exception (plan D2): the staff review path is outside the tenant boundary.
        pool = request.app.state.pool
        async with pool.acquire() as conn:
            # Admin review path reads the shared aa_internal reference pool directly
            # (published_tours.tenant_id sentinel) — unchanged, not a per-tenant lookup.
            pt = await conn.fetchrow("""
                SELECT * FROM gold_aa_internal.published_tours
                WHERE id = $1::uuid AND tenant_id = $2::uuid
            """, tour_id, tenant_id)
            if not pt:
                raise HTTPException(status_code=404, detail="Tour not found")
            raw = await conn.fetchrow(_TOUR_RAW_SQL, pt["tour_id"])
            # AA-796 — generated_content (A1 writer output incl. metadata llm_cost/model) and its
            # quality_scores row are admin review data; the tenant-JWT branch never reads them.
            gen = await conn.fetchrow("""
                SELECT * FROM silver_aa_internal.generated_content
                WHERE tour_id = $1::uuid
                ORDER BY version_num DESC LIMIT 1
            """, pt["tour_id"])
            qs = await conn.fetchrow("""
                SELECT * FROM silver_aa_internal.quality_scores
                WHERE generated_content_id = $1::uuid
                ORDER BY evaluated_at DESC LIMIT 1
            """, gen["id"]) if gen else None
            sc = await conn.fetchrow(_TOUR_SEO_SQL, pt["tour_id"])
        # admin tour-full/[id] review panel — unchanged
        return {
            "published": _safe_row(pt),
            "raw": _safe_row(raw),
            "generated": _safe_row(gen),
            "quality": _safe_row(qs),
            "seo": _safe_row(sc),
        }

    # Real tenant JWT: scope through tenant_tour_versions (AA-582 fix), same pattern as v1_marketplace.py —
    # published_tours.tenant_id is always the aa_internal sentinel and was never a valid per-tenant filter.
    # AA-859 D5b: on the tenant pool behind its flag.
    raw = sc = None
    async with acquire_scoped_conn(
        request, tenant_id, T2_VERSION_DETAIL_READS_FLAG, fail_closed=True,
    ) as (conn, used_tenant_pool):
        pt = await conn.fetchrow(_TOUR_VIA_VERSION_SQL.format(tenant_cond=_TOUR_VIA_VERSION_TENANT_COND),
                                 tour_id, tenant_id)
        if pt:
            raw = await conn.fetchrow(_TOUR_RAW_SQL, pt["tour_id"])
            sc = await conn.fetchrow(_TOUR_SEO_SQL, pt["tour_id"])
    _log_t2_detail_read("full", tenant_id, used_tenant_pool, found=pt is not None)
    if not pt:
        raise HTTPException(status_code=404, detail="Tour not found")

    # AA-796 — tenant JWT (reachable through the generic /api/tenant/[...path] proxy): content
    # only, no A1 score / review state / writer output.
    return {
        "published": _safe_row(project_tour(pt)),
        "raw": _safe_row(project_raw(raw) if raw else None),
        "seo": _safe_row(sc),
    }


class TourEditRequest(_BM):
    field: str
    value: str
    approved_by: Optional[str] = "content_team"


@router.get("/{tour_id}")
async def get_tour(
    tour_id: str,
    request: Request,
    tenant=Depends(get_tenant),
):
    tenant_id = tenant["sub"]

    # AA-582 fix: scope through tenant_tour_versions (v1_marketplace.py pattern) —
    # published_tours.tenant_id is always the aa_internal sentinel, never per-tenant.
    # AA-859 D5b: on the tenant pool behind T2_VERSION_DETAIL_READS_FLAG (same SQL as /full's tenant branch).
    async with acquire_scoped_conn(
        request, tenant_id, T2_VERSION_DETAIL_READS_FLAG, fail_closed=True,
    ) as (conn, used_tenant_pool):
        row = await conn.fetchrow(_TOUR_VIA_VERSION_SQL.format(tenant_cond=_TOUR_VIA_VERSION_TENANT_COND),
                                  tour_id, tenant_id)
    _log_t2_detail_read("tour", tenant_id, used_tenant_pool, found=row is not None)

    if not row:
        raise HTTPException(status_code=404, detail="Tour not found")

    return project_tour(row)  # AA-796 — no quality_score / quality_score_id


# ── P3-S4: Version Detail + Approve/Reject/Edit ───────────────────────────────

# AA-796 — the hidden version fields now live in api/routers/_tenant_projection.py.

# GET /v1/tours/versions/{version_id}. `{job_col}` / `{job_join}` as in _MY_VERSIONS_ROWS_SQL; on the tenant
# pool `job_status` is filled from `ttv.job_id` (part of ttv.*) on MAIN.
_VERSION_TENANT_COND = "ttv.id = $1::uuid AND ttv.tenant_id = $2::uuid"
_VERSION_SQL = """
    SELECT ttv.*, pt.tour_id, pt.aa_name, pt.aa_subtitle, pt.aa_summary,
           pt.aa_description, pt.aa_highlights, pt.aa_itineraries,
           pt.seo_title AS aa_seo_title, pt.seo_meta AS aa_seo_meta,
           pt.updated_at AS master_updated_at,  -- AA-990: master_newer is derived from it (project_version)
           rt.country, rt.duration, rt.price_raw,
           rt.inclusions, rt.exclusions,
           {job_col}  -- AA-652
    FROM gold_aa_internal.tenant_tour_versions ttv
    JOIN gold_aa_internal.published_tours pt ON pt.id = ttv.published_tour_id
    LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
    {job_join}
    WHERE {tenant_cond}
"""
_VERSION_JOB_JOINED = {"job_col": "j.status AS job_status", "job_join": _JOB_JOIN}
_VERSION_JOB_SPLIT = {"job_col": "NULL::text AS job_status", "job_join": ""}
_VERSION_HISTORY_TENANT_COND = "tenant_id = $1::uuid"
_VERSION_HISTORY_SQL = """
    SELECT id, version_number, status, edit_source, created_at
    FROM gold_aa_internal.tenant_tour_versions
    WHERE {tenant_cond} AND published_tour_id = $2::uuid
    ORDER BY version_number ASC
"""


@router.get("/versions/{version_id}")
async def get_version(
    version_id: str,
    request: Request,
    tenant=Depends(get_tenant),
):
    """Get version detail with AA original for before/after diff.

    AA-990: also `master_updated_at` and `master_newer` (see GET /v1/tours/my-versions)."""
    tenant_id = tenant["sub"]

    history = []
    async with acquire_scoped_conn(
        request, tenant_id, T2_VERSION_READS_FLAG, fail_closed=True,
    ) as (conn, used_tenant_pool):
        job = _VERSION_JOB_SPLIT if used_tenant_pool else _VERSION_JOB_JOINED
        row = await conn.fetchrow(_VERSION_SQL.format(tenant_cond=_VERSION_TENANT_COND, **job),
                                  version_id, tenant_id)
        served = {}
        if row:  # version history
            history = await conn.fetch(_VERSION_HISTORY_SQL.format(tenant_cond=_VERSION_HISTORY_TENANT_COND),
                                       tenant_id, row["published_tour_id"])
            served = (await _annotate_served(conn, [{"tour_id": row.get("tour_id")}]))[0]   # AA-986 L3
    _log_t2_read("version", tenant_id, used_tenant_pool, found=row is not None)

    if not row:
        raise HTTPException(status_code=404, detail="Version not found")
    if used_tenant_pool:  # after the tenant block: the tenant connection is released before MAIN
        row = (await _with_job_status(request.app.state.pool, [row], "job_id"))[0]

    # AA-770 (S233) — scores and QA state stay admin-side (/admin/content-quality).
    # AA-796 — and the score-derived needs_review status reads as ai_generated.
    return {
        **project_version(row),
        **{k: v for k, v in served.items() if k != "tour_id"},
        "version_history": [project_version(h) for h in history],
    }


# AA-566 Phần B.2 — tenant-facing DOCX export, one version at a time (matches Admin Master
# Content's own single-version DOCX convention, admin_pipeline.py::export_tour_version_docx()
# — same python-docx pattern, same off-white-background/heading-color helpers, deliberately NOT
# a shared function since the two read completely different tables (silver_aa_internal.
# generated_content there vs gold_aa_internal.tenant_tour_versions here) and the tenant version
# has no SEO/score/judge fields to render at all (AA-566 Phần B.5 — those are being dropped from
# the tenant-facing drawer for the same reason: internal jargon, not tenant content).
_EXPORT_VERSION_TENANT_COND = "ttv.tenant_id = $2::uuid"
_EXPORT_VERSION_SQL = """
            SELECT ttv.version_number, ttv.rewritten_content, ttv.rewrite_language,
                   pt.tour_id, pt.aa_name, pt.aa_subtitle, pt.aa_summary, pt.aa_highlights,
                   pt.aa_itineraries,
                   rt.country, rt.duration, rt.inclusions, rt.exclusions
            FROM gold_aa_internal.tenant_tour_versions ttv
            JOIN gold_aa_internal.published_tours pt ON pt.id = ttv.published_tour_id
            LEFT JOIN silver_aa_internal.raw_tours rt ON rt.tour_id = pt.tour_id
            WHERE ttv.id = $1::uuid AND {tenant_cond}
        """


@router.get("/versions/{version_id}/export-docx")
async def export_version_docx(
    version_id: str,
    request: Request,
    tenant=Depends(get_tenant),
):
    import json as _json
    tenant_id = tenant["sub"]
    pool = request.app.state.pool

    # AA-859 D5b: on the tenant pool behind T2_VERSION_DETAIL_READS_FLAG (229 on tenant_tour_versions, 239
    # master read on published_tours / raw_tours).
    async with acquire_scoped_conn(
        request, tenant_id, T2_VERSION_DETAIL_READS_FLAG, fail_closed=True,
    ) as (conn, used_tenant_pool):
        row = await conn.fetchrow(_EXPORT_VERSION_SQL.format(tenant_cond=_EXPORT_VERSION_TENANT_COND),
                                  version_id, tenant_id)
    _log_t2_detail_read("export_docx", tenant_id, used_tenant_pool, found=row is not None)
    if not row:
        raise HTTPException(status_code=404, detail="Version not found")

    # AA-802 — record the download so the admin Content Quality drawer can tell an admin that the
    # tenant already holds a copy older than a later admin correction. Admin-only row (tenants never
    # read audit_log). Best-effort, like the T9 export audit: a lost row must not fail the download.
    # AA-859 MAIN exception: audit_log has no RLS and no aa_app_user grant; written on the admin pool after
    # the tenant block released its connection.
    try:
        import uuid as _uuid
        await write_audit_log(
            pool, tenant_id=str(tenant_id), actor=f"tenant:{tenant_id}",
            action=TenantAuditAction.TOUR_VERSION_EXPORTED, resource_type="tenant_tour_version",
            # canonical uuid text: the admin drawer matches it against ttv.id::text
            resource_id=str(_uuid.UUID(str(version_id))),
            details={"format": "docx", "version_number": row["version_number"]},
        )
    except Exception:
        pass

    # Same rewritten_content-overrides-aa_* merge CatalogTab.tsx's own loadDetail() does —
    # a tenant's manual edit lives in rewritten_content, not the flat aa_* columns.
    rc = {}
    if row["rewritten_content"]:
        try:
            raw_rc = row["rewritten_content"]
            rc = _json.loads(raw_rc) if isinstance(raw_rc, str) else dict(raw_rc)
        except Exception:
            rc = {}
    name = rc.get("name") or row["aa_name"] or "(untitled)"
    subtitle = rc.get("subtitle") or row["aa_subtitle"] or ""
    summary = rc.get("summary") or row["aa_summary"] or ""
    highlights = rc.get("highlights")
    if not isinstance(highlights, list):
        try:
            highlights = _json.loads(row["aa_highlights"]) if row["aa_highlights"] else []
        except Exception:
            highlights = []
    itineraries = rc.get("itineraries") or row["aa_itineraries"] or ""

    import io
    from docx import Document
    from docx.shared import Pt, RGBColor
    from fastapi.responses import StreamingResponse

    ORANGE = RGBColor(0xDB, 0x96, 0x28)
    BLACKBLUE = RGBColor(0x1F, 0x29, 0x33)
    GRAY = RGBColor(0x33, 0x36, 0x3D)

    doc = Document()

    def _section(text):
        p = doc.add_paragraph()
        bar = p.add_run("▌ ")
        bar.bold = True
        bar.font.color.rgb = ORANGE
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = BLACKBLUE

    def _kv(label, value):
        p = doc.add_paragraph()
        lr = p.add_run(f"{label}: ")
        lr.bold = True
        lr.font.size = Pt(10.5)
        lr.font.color.rgb = BLACKBLUE
        vr = p.add_run("" if value is None else str(value))
        vr.font.size = Pt(10.5)
        vr.font.color.rgb = GRAY

    def _body(text):
        p = doc.add_paragraph()
        for i, line in enumerate(str(text or "").split("\n")):
            if i > 0:
                p.add_run().add_break()
            r = p.add_run(line)
            r.font.size = Pt(10.5)
            r.font.color.rgb = GRAY

    title = doc.add_paragraph()
    tr = title.add_run(name)
    tr.bold = True
    tr.font.size = Pt(20)
    tr.font.color.rgb = BLACKBLUE
    if subtitle:
        sub = doc.add_paragraph()
        sr = sub.add_run(subtitle)
        sr.italic = True
        sr.font.size = Pt(12)
        sr.font.color.rgb = ORANGE
    meta = doc.add_paragraph()
    mr = meta.add_run(f"Version {row['version_number']}   ·   {row['rewrite_language']}")
    mr.font.size = Pt(10)
    mr.font.color.rgb = ORANGE

    _section("Trip Details")
    _kv("Country", row["country"] or "—")
    _kv("Duration", row["duration"] or "—")

    _section("Summary")
    _body(summary)

    if highlights:
        _section("Highlights")
        for h in highlights:
            doc.add_paragraph(str(h), style="List Bullet")

    if itineraries:
        _section("Itinerary")
        _body(itineraries)

    if row["inclusions"]:
        _section("Inclusions")
        _body(row["inclusions"])

    if row["exclusions"]:
        _section("Exclusions")
        _body(row["exclusions"])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    short = str(row["tour_id"]).split("-")[0]
    return StreamingResponse(
        iter([buf.read()]),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=tour_{short}_v{row['version_number']}.docx"},
    )


class VersionActionRequest(_BM):
    action: str  # approve | reject | edit
    edited_content: Optional[dict] = None
    edited_by: Optional[str] = None


# ── AA-859 W6: PATCH /v1/tours/versions/{id} on the tenant pool (T2_VERSION_WRITE_FLAG) ──────────────
# Flag off: the admin pool, exactly the old statements (no transaction, no lock). Flag on: one
# transaction on the aa_app_user tenant pool (RLS on tenant_tour_versions, migration 229; grants SELECT,
# INSERT, UPDATE), fail closed:
#   1. own-row SELECT -> 404;
#   2. edit: advisory xact lock per (tenant, published tour), MAX(version_number) + 1, INSERT. No
#      FOR UPDATE: PostgreSQL rejects locking clauses with aggregates (0A000). Two-int lock form, so the
#      key space cannot meet the one-bigint locks used elsewhere (t2_rewrite, slate, publish limit, ...);
#      it needs no grant. A duplicate number anyway (UNIQUE (tenant_id, published_tour_id,
#      version_number), e.g. a concurrent rewrite on MAIN) -> 409 version_conflict, never a 500;
#   3. approve / reject: status UPDATE.
# No audit row, no enqueue, no external call: nothing here runs on MAIN after the transaction.
# SQL in module constants so tests can mutate them (forgotten-WHERE, lock removed).
AA859_TTV_LOCK_CLASS = 859
VERSION_CONFLICT_DETAIL = "version_conflict: another edit of this tour was saved at the same time, reload and retry"

_VERSION_OWN_ROW_SQL = """
            SELECT id, published_tour_id, version_number, rewritten_content
            FROM gold_aa_internal.tenant_tour_versions
            WHERE id = $1::uuid AND tenant_id = $2::uuid
        """
_VERSION_EDIT_LOCK_SQL = "SELECT pg_advisory_xact_lock($1::int, hashtext($2::text || ':' || $3::text))"
_VERSION_NEXT_NUMBER_SQL = """
                SELECT COALESCE(MAX(version_number), 0) + 1
                FROM gold_aa_internal.tenant_tour_versions
                WHERE tenant_id = $1::uuid AND published_tour_id = $2::uuid
            """
_VERSION_EDIT_INSERT_SQL = """
                INSERT INTO gold_aa_internal.tenant_tour_versions
                    (tenant_id, published_tour_id, version_number,
                     parent_version_id, rewritten_content, status,
                     edit_source, edited_at, edited_by_user)
                VALUES ($1::uuid, $2::uuid, $3,
                        $4::uuid, $5::jsonb, 'pending',
                        'tenant_edit', NOW(), $6)
                RETURNING id
            """
_VERSION_STATUS_UPDATE_SQL = """
                UPDATE gold_aa_internal.tenant_tour_versions
                SET status = $1, edited_at = NOW()
                WHERE id = $2::uuid AND tenant_id = $3::uuid
            """


@router.patch("/versions/{version_id}")
async def update_version(
    version_id: str,
    body: VersionActionRequest,
    request: Request,
    tenant=Depends(get_tenant),
):
    """Approve, reject, or inline-edit a tenant tour version (AA-859 W6: tenant pool behind its flag)."""
    tenant_id = tenant["sub"]

    if body.action not in ("approve", "reject", "edit"):
        raise HTTPException(status_code=400, detail="action must be approve|reject|edit")
    # AA-859 PR-D6 review: an edit without content used to fall through to the status branch and silently
    # REJECT the caller's version (200 {"status": "rejected"}). Both paths, before any acquire.
    if body.action == "edit" and not body.edited_content:
        raise HTTPException(status_code=400, detail="edited_content is required for action=edit")

    async with acquire_scoped_conn(
        request, tenant_id, T2_VERSION_WRITE_FLAG, fail_closed=True,
    ) as (conn, used_tenant_pool):
        logger.info("aa859_t2_version_write", tenant_id=str(tenant_id), version_id=version_id,
                    action=body.action, used_tenant_pool=used_tenant_pool)
        return await _apply_version_action(conn, tenant_id, version_id, body, guarded=used_tenant_pool)


async def _apply_version_action(conn, tenant_id, version_id: str, body: VersionActionRequest, *,
                                guarded: bool) -> dict:
    """The route's statements on `conn`. guarded=True (tenant pool, inside acquire_scoped_conn's
    transaction): the edit takes the advisory lock and a duplicate version number is a 409. guarded=False
    (admin pool, flag off): the pre-AA-859 statements, unchanged."""
    import json as _json
    row = await conn.fetchrow(_VERSION_OWN_ROW_SQL, version_id, tenant_id)
    if not row:
        raise HTTPException(status_code=404, detail="Version not found")

    if body.action == "edit":
        # Inline edit (edited_content checked non-empty before the acquire) → create new version
        if guarded:
            await conn.execute(_VERSION_EDIT_LOCK_SQL, AA859_TTV_LOCK_CLASS, str(tenant_id),
                               str(row["published_tour_id"]))
        next_ver = await conn.fetchval(_VERSION_NEXT_NUMBER_SQL, tenant_id, row["published_tour_id"])
        args = (tenant_id, row["published_tour_id"], next_ver, version_id,
                _json.dumps(body.edited_content), body.edited_by or "tenant")
        if guarded:
            try:
                new_id = await conn.fetchval(_VERSION_EDIT_INSERT_SQL, *args)
            except asyncpg.UniqueViolationError as exc:
                logger.warning("aa859_t2_version_conflict", tenant_id=str(tenant_id), version_id=version_id,
                               version_number=next_ver)
                raise HTTPException(status_code=409, detail=VERSION_CONFLICT_DETAIL) from exc
        else:
            new_id = await conn.fetchval(_VERSION_EDIT_INSERT_SQL, *args)
        return {
            "status": "edited",
            "new_version_id": str(new_id),
            "version_number": next_ver,
        }

    new_status = "approved" if body.action == "approve" else "rejected"
    await conn.execute(_VERSION_STATUS_UPDATE_SQL, new_status, version_id, tenant_id)
    return {"status": new_status, "version_id": version_id}
# P3 complete Tue May  5 11:42:07 +07 2026


# AA-859 (PR-D6): PATCH /v1/tours/{tour_id}/approve removed. No caller in frontend/ (its last one,
# (internal)/catalog/page.tsx, went with AA-663), and it could never work: it updated
# gold_aa_internal.published_tours WHERE tenant_id = <tenant>, 0 rows for every tenant (published_tours
# is all AA master content), and still answered ok: true.
